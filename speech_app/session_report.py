"""Rapport de fin de session : qui parle, traduction relue, résumé daté.

Entrée : une `Session` (sessions.py) avec audio.wav, segments.jsonl {id, debut_s, fin_s, ar, fr}.
Sorties dans `session.dossier_rapport` : transcript.md, resume.md, sous_titres.srt, rapport.json
(et transcript.docx si python-docx est installé).

Quatre étapes, chacune testable seule et remplaçable par un faux (options["diariseur"], options["llm"]) :

  a. diarisation   pyannote/speaker-diarization-community-1 (HF_TOKEN, modèle à accès restreint) si possible ;
                   sinon repli libre : embeddings ECAPA (speechbrain) par segment + regroupement agglomératif.
  b. révision      un LLM local relit la traduction segment par segment, avec les voisins et le lexique
                   donnees/noms/noms_v1.csv (noms propres). L'original est gardé à côté.
  c. résumé        en français, structuré, chaque point daté [mm:ss] ; sans LLM : résumé extractif signalé.
  d. écriture      transcript.md, resume.md, sous_titres.srt, rapport.json.

LLM : variable PORTEPAROLE_LLM_RAPPORT (défaut Qwen/Qwen2.5-7B-Instruct en 4 bits nf4, calcul fp16, adapté au T4 ;
« off » le désactive). Il est chargé APRÈS avoir libéré le groupe de modèles de l'interface, sous le verrou GPU
global (lazy_engines.GPU_LOCK), puis déchargé à la fin.
"""
from __future__ import annotations

import csv
import gc
import json
import os
import re
import threading
import time
import unicodedata
import wave
from collections import Counter
from pathlib import Path
from typing import Callable

import numpy as np

from sessions import FREQUENCE, Session

LLM_DEFAUT = "Qwen/Qwen2.5-7B-Instruct"
LLM_REPLI = "Qwen/Qwen2.5-1.5B-Instruct"          # déjà utilisé par S3
DIARISATION_PYANNOTE = "pyannote/speaker-diarization-community-1"
ECAPA = "speechbrain/spkrec-ecapa-voxceleb"
LEXIQUE_DEFAUT = Path(__file__).resolve().parent.parent / "donnees" / "noms" / "noms_v1.csv"
FICHIERS = ("transcript.md", "resume.md", "sous_titres.srt", "rapport.json", "transcript.docx")

OPTIONS_DEFAUT = {
    "diarisation": "auto",        # auto | pyannote | ecapa | off
    "nb_locuteurs": None,         # connu à l'avance ? sinon estimé
    "seuil_similarite": 0.40,     # ECAPA : deux segments de même voix ont une similarité cosinus au-dessus
    "revision": True,             # relecture par LLM
    "resume": True,
    "inclure_arabe": False,       # arabe dans transcript.md
    "docx": True,
    "contexte_avant": 2,
    "contexte_apres": 1,
}

# ---------------------------------------------------------------- verrou GPU et mémoire
_VERROU_LOCAL = threading.RLock()


def verrou_gpu():
    """Verrou GPU global de lazy_engines (créé par l'autre module) ; verrou local si absent."""
    try:
        import lazy_engines
        return getattr(lazy_engines, "GPU_LOCK", None) or _VERROU_LOCAL
    except Exception:                                   # noqa: BLE001 : import tolérant (torch absent, etc.)
        return _VERROU_LOCAL


def liberer_gpu(pool=None) -> None:
    """Décharge le groupe de modèles de l'interface (Pool) puis vide le cache CUDA."""
    if pool is not None:
        if hasattr(pool, "decharger"):
            pool.decharger()
        else:                                           # Pool actuel : pas de méthode publique
            try:
                pool._engines, pool._group = {}, None
            except Exception:                           # noqa: BLE001
                pass
    gc.collect()
    try:
        import torch
        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    except ImportError:
        pass


def _cuda() -> bool:
    try:
        import torch
        return torch.cuda.is_available()
    except ImportError:
        return False


# ---------------------------------------------------------------- utilitaires
def horodatage(s: float) -> str:
    s = int(max(0, s))
    h, m, sec = s // 3600, s % 3600 // 60, s % 60
    return f"{h}:{m:02d}:{sec:02d}" if h else f"{m:02d}:{sec:02d}"


def _ts_srt(t: float) -> str:
    ms = int(round(max(0.0, t) * 1000))
    return f"{ms // 3600000:02d}:{ms % 3600000 // 60000:02d}:{ms % 60000 // 1000:02d},{ms % 1000:03d}"


def lire_audio(chemin: Path) -> np.ndarray:
    """wav -> float32 mono 16 kHz."""
    try:
        with wave.open(str(chemin), "rb") as f:
            if f.getsampwidth() == 2 and f.getframerate() == FREQUENCE:
                x = np.frombuffer(f.readframes(f.getnframes()), dtype="<i2").astype(np.float32) / 32768.0
                return x.reshape(-1, f.getnchannels()).mean(axis=1) if f.getnchannels() > 1 else x
    except wave.Error:
        pass
    import librosa
    return librosa.load(str(chemin), sr=FREQUENCE, mono=True)[0]


def texte_de(seg: dict) -> str:
    return (seg.get("fr") or "").strip()


# ---------------------------------------------------------------- lexique de noms propres
_DIACRITIQUES = re.compile("[ً-ْـ]")
_LETTRES_AR = "ء-ي"


def normaliser_ar(s: str) -> str:
    s = _DIACRITIQUES.sub("", unicodedata.normalize("NFKC", s or ""))
    for a, b in (("أ", "ا"), ("إ", "ا"), ("آ", "ا"), ("ى", "ي"), ("ة", "ه")):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()


def charger_lexique(chemin: str | Path | None = None) -> list[dict]:
    """Lit noms_v1.csv : forme_arabe, variantes (séparées par « | »), forme_francaise, type, id_wikidata."""
    chemin = Path(chemin or os.environ.get("PORTEPAROLE_LEXIQUE") or LEXIQUE_DEFAUT)
    if not chemin.is_file():
        return []
    out = []
    with open(chemin, encoding="utf-8", newline="") as f:
        for ligne in csv.DictReader(f):
            fr, ar = (ligne.get("forme_francaise") or "").strip(), (ligne.get("forme_arabe") or "").strip()
            if not fr or not ar:
                continue
            formes = [ar] + [v.strip() for v in (ligne.get("variantes") or "").split("|") if v.strip()]
            formes = sorted({normaliser_ar(x) for x in formes if len(normaliser_ar(x)) >= 4})
            if formes:
                out.append({"fr": fr, "type": (ligne.get("type") or "").strip(), "id": ligne.get("id_wikidata"),
                            "formes": [re.compile(rf"(?<![{_LETTRES_AR}])[وفبلك]?(?:ال)?{re.escape(x)}(?![{_LETTRES_AR}])")
                                       for x in formes]})
    return out


def entites_du_texte(ar: str, lexique: list[dict]) -> list[dict]:
    """Entités du lexique dont une forme arabe apparaît dans `ar`."""
    n = normaliser_ar(ar)
    if not n:
        return []
    return [{"fr": e["fr"], "type": e["type"], "id": e["id"]} for e in lexique if any(p.search(n) for p in e["formes"])]


# ---------------------------------------------------------------- a. diarisation
def regrouper_embeddings(emb: np.ndarray, seuil_similarite: float = 0.40, nb_locuteurs: int | None = None) -> list[int]:
    """Regroupement agglomératif (lien moyen, distance cosinus). S'arrête quand la meilleure paire a une
    similarité < seuil, ou quand il reste `nb_locuteurs` groupes. Renvoie une étiquette par ligne."""
    n = len(emb)
    if n == 0:
        return []
    e = emb / np.maximum(np.linalg.norm(emb, axis=1, keepdims=True), 1e-9)
    sim = e @ e.T
    groupes = [[i] for i in range(n)]
    cible = max(1, nb_locuteurs or 1)
    while len(groupes) > cible:
        meilleur, paire = -2.0, None
        for a in range(len(groupes)):
            for b in range(a + 1, len(groupes)):
                s = float(sim[np.ix_(groupes[a], groupes[b])].mean())
                if s > meilleur:
                    meilleur, paire = s, (a, b)
        if paire is None or (nb_locuteurs is None and meilleur < seuil_similarite):
            break
        a, b = paire
        groupes[a] += groupes[b]
        del groupes[b]
    etiquettes = [0] * n
    for k, g in enumerate(groupes):
        for i in g:
            etiquettes[i] = k
    return etiquettes


def attribuer_locuteurs(segments: list[dict], tours: list[tuple[float, float, str]]) -> list[str | None]:
    """Locuteur majoritaire (plus long recouvrement) sur l'intervalle de chaque segment ; à défaut, le tour le plus proche."""
    out: list[str | None] = []
    for seg in segments:
        d, f = float(seg.get("debut_s") or 0), float(seg.get("fin_s") or 0)
        temps: Counter = Counter()
        for a, b, spk in tours:
            rec = min(f, b) - max(d, a)
            if rec > 0:
                temps[spk] += rec
        if temps:
            out.append(temps.most_common(1)[0][0])
        elif tours:
            milieu = (d + f) / 2
            out.append(min(tours, key=lambda t: min(abs(milieu - t[0]), abs(milieu - t[1])))[2])
        else:
            out.append(None)
    return out


def _numeroter(etiquettes: list) -> list[int | None]:
    """Étiquettes quelconques -> 1, 2, 3… dans l'ordre d'apparition."""
    ordre: dict = {}
    for e in etiquettes:
        if e is not None and e not in ordre:
            ordre[e] = len(ordre) + 1
    return [ordre.get(e) for e in etiquettes]


def diariser_pyannote(chemin_audio: Path, nb_locuteurs: int | None = None) -> list[tuple[float, float, str]]:
    """Tours de parole via pyannote.audio 4 (community-1). Lève une exception si indisponible."""
    token = os.environ.get("HF_TOKEN") or os.environ.get("HUGGING_FACE_HUB_TOKEN")
    if not token:
        raise RuntimeError("HF_TOKEN absent : pyannote/speaker-diarization-community-1 demande un accès accepté")
    import torch
    from pyannote.audio import Pipeline
    try:
        pipeline = Pipeline.from_pretrained(DIARISATION_PYANNOTE, token=token)
    except TypeError:                                    # pyannote 3.x
        pipeline = Pipeline.from_pretrained(DIARISATION_PYANNOTE, use_auth_token=token)
    if pipeline is None:
        raise RuntimeError("accès au modèle pyannote refusé (conditions non acceptées ?)")
    if _cuda():
        pipeline.to(torch.device("cuda"))
    forme_onde = torch.from_numpy(lire_audio(chemin_audio)).unsqueeze(0)     # évite torchcodec/ffmpeg
    kw = {"num_speakers": nb_locuteurs} if nb_locuteurs else {}
    sortie = pipeline({"waveform": forme_onde, "sample_rate": FREQUENCE}, **kw)
    annotation = getattr(sortie, "speaker_diarization", sortie)
    tours = [(float(t.start), float(t.end), str(spk)) for t, _, spk in annotation.itertracks(yield_label=True)]
    del pipeline
    liberer_gpu()
    return tours


def diariser_ecapa(chemin_audio: Path, segments: list[dict], seuil_similarite: float = 0.40,
                   nb_locuteurs: int | None = None) -> list[int]:
    """Repli libre : un embedding ECAPA par segment, regroupement agglomératif. Une étiquette par segment."""
    import torch
    try:
        from speechbrain.inference.speaker import EncoderClassifier
    except ImportError:                                  # speechbrain < 1.0
        from speechbrain.pretrained import EncoderClassifier
    signal = lire_audio(chemin_audio)
    dossier = Path(os.environ.get("PORTEPAROLE_CACHE", Path(__file__).resolve().parent / "cache")) / "ecapa"
    clf = EncoderClassifier.from_hparams(source=ECAPA, savedir=str(dossier),
                                         run_opts={"device": "cuda" if _cuda() else "cpu"})
    embs = []
    for seg in segments:
        a, b = float(seg.get("debut_s") or 0), float(seg.get("fin_s") or 0)
        if b - a < 1.0:                                   # trop court pour une voix : on élargit autour du centre
            m = (a + b) / 2
            a, b = max(0.0, m - 0.5), m + 0.5
        morceau = signal[int(a * FREQUENCE):int(b * FREQUENCE)]
        if len(morceau) < FREQUENCE // 4:
            morceau = np.pad(morceau, (0, FREQUENCE // 4 - len(morceau)))
        with torch.inference_mode():
            v = clf.encode_batch(torch.from_numpy(morceau).float().unsqueeze(0).to(clf.device))
        embs.append(v.squeeze().detach().cpu().numpy())
    del clf
    liberer_gpu()
    return regrouper_embeddings(np.stack(embs), seuil_similarite, nb_locuteurs) if embs else []


def etape_diarisation(session: Session, segments: list[dict], options: dict, avertissements: list[str]) -> tuple[list[int | None], str]:
    """Renvoie (locuteur 1..N par segment, nom de la méthode)."""
    if options.get("diariseur"):                         # faux ou moteur externe : callable -> (étiquettes, méthode)
        etiquettes, methode = options["diariseur"](session, segments, options)
        return _numeroter(list(etiquettes)), methode
    mode = options.get("diarisation", "auto")
    if mode == "off" or not segments:
        return [1] * len(segments), "aucune (un seul locuteur supposé)"
    if not session.chemin_audio.is_file():
        avertissements.append("audio.wav absent : diarisation impossible")
        return [1] * len(segments), "aucune (audio absent)"
    if mode in ("auto", "pyannote"):
        try:
            tours = diariser_pyannote(session.chemin_audio, options.get("nb_locuteurs"))
            return _numeroter(attribuer_locuteurs(segments, tours)), f"pyannote ({DIARISATION_PYANNOTE})"
        except Exception as exc:                         # noqa: BLE001 : repli sur ECAPA, cause gardée
            avertissements.append(f"pyannote indisponible ({type(exc).__name__}: {str(exc)[:150]}) : repli ECAPA")
    try:
        et = diariser_ecapa(session.chemin_audio, segments, options.get("seuil_similarite", 0.40), options.get("nb_locuteurs"))
        return _numeroter(et), f"ECAPA speechbrain + regroupement agglomératif (seuil {options.get('seuil_similarite', 0.40)})"
    except Exception as exc:                             # noqa: BLE001
        avertissements.append(f"diarisation ECAPA impossible ({type(exc).__name__}: {str(exc)[:150]}) : un seul locuteur supposé")
        return [1] * len(segments), "aucune (échec des deux méthodes)"


# ---------------------------------------------------------------- LLM local
class LLMLocal:
    """Qwen2.5 en 4 bits nf4 (calcul fp16) sur GPU ; appel : llm(messages, max_new_tokens) -> texte."""

    def __init__(self, model_id: str | None = None):
        self.model_id = model_id or os.environ.get("PORTEPAROLE_LLM_RAPPORT") or LLM_DEFAUT
        self.modele = self.tokenizer = None
        self.nom = self.model_id

    def charger(self) -> "LLMLocal":
        if self.modele is not None:
            return self
        import torch
        from transformers import AutoModelForCausalLM, AutoTokenizer
        if not torch.cuda.is_available():
            raise RuntimeError("pas de GPU : le LLM de relecture n'est pas chargé")
        try:
            from transformers import BitsAndBytesConfig
            import bitsandbytes  # noqa: F401
            q = BitsAndBytesConfig(load_in_4bit=True, bnb_4bit_quant_type="nf4",
                                   bnb_4bit_compute_dtype=torch.float16, bnb_4bit_use_double_quant=True)
            self.tokenizer = AutoTokenizer.from_pretrained(self.model_id)
            self.modele = AutoModelForCausalLM.from_pretrained(self.model_id, quantization_config=q, device_map={"": 0})
            self.nom = f"{self.model_id} (4 bits nf4)"
        except Exception:                                # noqa: BLE001 : bitsandbytes absent ou échec -> petit modèle fp16
            self.decharger()
            self.model_id = LLM_REPLI
            self.tokenizer = AutoTokenizer.from_pretrained(self.model_id)
            try:
                self.modele = AutoModelForCausalLM.from_pretrained(self.model_id, dtype=torch.float16).to("cuda")
            except TypeError:
                self.modele = AutoModelForCausalLM.from_pretrained(self.model_id, torch_dtype=torch.float16).to("cuda")
            self.nom = f"{self.model_id} (fp16, repli)"
        self.modele.eval()
        return self

    def __call__(self, messages: list[dict], max_new_tokens: int = 256) -> str:
        import torch
        self.charger()
        texte = self.tokenizer.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        entree = self.tokenizer(texte, return_tensors="pt").to(self.modele.device)
        with torch.inference_mode():
            sortie = self.modele.generate(**entree, max_new_tokens=max_new_tokens, do_sample=False,
                                          repetition_penalty=1.05, pad_token_id=self.tokenizer.eos_token_id)
        return self.tokenizer.decode(sortie[0][entree["input_ids"].shape[1]:], skip_special_tokens=True).strip()

    def decharger(self) -> None:
        self.modele = self.tokenizer = None
        liberer_gpu()


# ---------------------------------------------------------------- b. révision
SYSTEME_REVISION = (
    "Tu relis la traduction française d'un journal télévisé arabe. Corrige uniquement les fautes évidentes "
    "(mots absurdes, phrase cassée, ponctuation) et mets les noms propres sous la forme du GLOSSAIRE. "
    "N'ajoute aucune information, ne supprime aucune information, ne change pas les chiffres. "
    "Si la traduction est correcte, recopie-la telle quelle. Réponds par la phrase relue, rien d'autre.")


def prompt_revision(segments: list[dict], i: int, revises: list[str], entites: list[dict],
                    avant: int = 2, apres: int = 1) -> list[dict]:
    seg = segments[i]
    ctx_av = " / ".join(revises[max(0, i - avant):i]) or "(début)"
    ctx_ap = " / ".join(texte_de(s) for s in segments[i + 1:i + 1 + apres]) or "(fin)"
    gloss = "\n".join(f"- {e['fr']} ({e['type']})" for e in entites) or "(aucun nom du glossaire dans ce segment)"
    user = (f"Contexte avant : {ctx_av}\nContexte après : {ctx_ap}\n\nGLOSSAIRE :\n{gloss}\n\n"
            f"Arabe : {seg.get('ar') or '(non disponible)'}\nTraduction à relire : {texte_de(seg)}")
    return [{"role": "system", "content": SYSTEME_REVISION}, {"role": "user", "content": user}]


def _nettoyer_sortie(s: str) -> str:
    s = s.strip().splitlines()[0].strip() if s.strip() else ""
    s = re.sub(r"^(Traduction( relue)?|Réponse)\s*:\s*", "", s, flags=re.I)
    return s.strip().strip('"«» ').strip()


def sortie_acceptable(original: str, revise: str) -> bool:
    """Garde-fou contre l'invention : longueur proche, pas d'arabe, mêmes chiffres."""
    if not revise or not original:
        return False
    r = len(revise) / max(1, len(original))
    if r > 1.6 or r < 0.5:
        return False
    if re.search(f"[{_LETTRES_AR}]", revise) and not re.search(f"[{_LETTRES_AR}]", original):
        return False
    return Counter(re.findall(r"\d+", original)) == Counter(re.findall(r"\d+", revise))


def etape_revision(segments: list[dict], llm: Callable, lexique: list[dict], options: dict,
                   progress: Callable | None = None) -> tuple[list[dict], dict]:
    """Renvoie (un dict par segment : fr_original, fr_revise, modifie, entites, raison), bilan du lexique."""
    revises: list[str] = []
    out = []
    appliques = {}
    for i, seg in enumerate(segments):
        original = texte_de(seg)
        ent = entites_du_texte(seg.get("ar") or "", lexique)
        for e in ent:
            appliques[e["fr"]] = e
        raison, final = "vide", original
        if original:
            try:
                brut = _nettoyer_sortie(llm(prompt_revision(segments, i, revises, ent, options.get("contexte_avant", 2),
                                                            options.get("contexte_apres", 1)),
                                            max_new_tokens=int(len(original) * 0.8) + 40))
            except Exception as exc:                     # noqa: BLE001 : un segment raté garde son original
                brut, raison = "", f"erreur LLM : {type(exc).__name__}"
            if brut and sortie_acceptable(original, brut):
                final, raison = brut, ("modifie" if brut != original else "inchange")
            elif brut:
                raison = "rejete (garde-fou) : original conservé"
        manquants = [e["fr"] for e in ent if e["fr"].lower() not in final.lower()]
        out.append({"fr_original": original, "fr_revise": final, "modifie": final != original, "raison": raison,
                    "entites": [e["fr"] for e in ent], "entites_absentes_du_texte": manquants})
        revises.append(final)
        if progress:
            progress((i + 1) / len(segments), f"Relecture {i + 1}/{len(segments)}")
    bilan = {"entites_detectees": sorted(appliques), "n_entites": len(appliques),
             "a_verifier": sorted({m for o in out for m in o["entites_absentes_du_texte"]})}
    return out, bilan


# ---------------------------------------------------------------- c. résumé
SYSTEME_RESUME = (
    "Tu es journaliste francophone. À partir du seul TRANSCRIPT fourni (traduction d'un journal arabe), écris un "
    "résumé en français. N'utilise AUCUNE information extérieure, n'invente rien : si une information manque, "
    "écris « non précisé ». Chaque point clé se termine par l'horodatage [mm:ss] du segment qui le fonde, "
    "recopié tel quel depuis le transcript. Format exact :\n"
    "## Sujet\nune phrase\n\n## Points clés\n- 3 à 6 puces, chacune finissant par [mm:ss]\n\n"
    "## Personnes, organisations, lieux cités\n- noms tels qu'ils figurent dans le transcript\n\n"
    "## Chiffres\n- chaque chiffre avec son contexte et [mm:ss] (ou « aucun chiffre cité »)")
MORCEAU_MAX = 6000                                       # caractères de transcript par appel (carte-réduction au-delà)


def lignes_transcript(segments: list[dict], textes: list[str], locuteurs: list[int | None]) -> list[str]:
    return [f"[{horodatage(s.get('debut_s') or 0)}] Locuteur {l or '?'} : {t}" for s, t, l in zip(segments, textes, locuteurs) if t]


def _decouper(lignes: list[str], taille: int) -> list[list[str]]:
    blocs, cur, n = [], [], 0
    for l in lignes:
        if cur and n + len(l) > taille:
            blocs.append(cur)
            cur, n = [], 0
        cur.append(l)
        n += len(l) + 1
    return blocs + ([cur] if cur else [])


def verifier_horodatages(markdown: str, segments: list[dict]) -> dict:
    valides = {horodatage(s.get("debut_s") or 0) for s in segments}
    cites = re.findall(r"\[(\d{1,3}:\d{2}(?::\d{2})?)\]", markdown)
    return {"cites": len(cites), "invalides": sorted({c for c in cites if c not in valides})}


def resume_llm(lignes: list[str], llm: Callable) -> str:
    blocs = _decouper(lignes, MORCEAU_MAX)
    if len(blocs) > 1:                                   # carte : notes par morceau, puis réduction
        notes = []
        for b in blocs:
            notes.append(llm([{"role": "system", "content": "Résume ce morceau de transcript en puces factuelles, "
                               "en gardant les noms, les chiffres et l'horodatage [mm:ss] de chaque fait. "
                               "N'ajoute rien qui ne figure pas dans le texte."},
                              {"role": "user", "content": "\n".join(b)}], max_new_tokens=500))
        corps = "NOTES PAR MORCEAU (extraites du transcript) :\n" + "\n\n".join(notes)
    else:
        corps = "TRANSCRIPT :\n" + "\n".join(blocs[0] if blocs else [])
    return llm([{"role": "system", "content": SYSTEME_RESUME}, {"role": "user", "content": corps}], max_new_tokens=900).strip()


_MOTS_VIDES = set("""le la les un une des de du et en à au aux ce cet cette ces il elle ils elles on nous vous je tu qui que quoi dont
où pour par sur dans avec sans sous est sont été être a ont avoir fait faire plus moins très aussi mais ou donc car se sa son ses leur leurs
y ne pas comme entre vers chez cela ceci celui celle ceux selon après avant""".split())


def resume_extractif(segments: list[dict], textes: list[str], lexique_ent: list[str], n_points: int = 5) -> str:
    """Résumé de secours sans modèle : segments les plus représentatifs, chiffres et noms propres repérés."""
    cand = [(i, t) for i, t in enumerate(textes) if len(t.split()) >= 4]
    freq = Counter(m for _, t in cand for m in set(re.findall(r"\w{4,}", t.lower())) if m not in _MOTS_VIDES)
    notes = []
    for i, t in cand:
        mots = [m for m in re.findall(r"\w{4,}", t.lower()) if m not in _MOTS_VIDES]
        notes.append((sum(freq[m] for m in set(mots)) / (1 + len(mots)) ** 0.5, i, t))
    choisis = sorted(sorted(notes, reverse=True)[:n_points], key=lambda x: x[1])
    ts = lambda i: horodatage(segments[i].get("debut_s") or 0)       # noqa: E731
    sujet = (cand[0][1] if cand else "non précisé")
    noms: dict[str, int] = {}
    for i, t in enumerate(textes):
        for m in re.finditer(r"(?<![.!?«]\s)(?<!^)\b([A-ZÉÈÀÂÔ][\wÀ-ÿ'’-]+(?:\s[A-ZÉÈÀÂÔ][\wÀ-ÿ'’-]+)*)", t):
            noms.setdefault(m.group(1), i)
    for e in lexique_ent:
        noms.setdefault(e, 0)
    chiffres = []
    for i, t in enumerate(textes):
        for m in re.finditer(r"\d+(?:[ .,]\d+)*(?:\s?(?:%|pour cent|millions?|milliards?|mille|dollars?|euros?))?", t):
            chiffres.append(f"- {m.group(0).strip()} : « {t[max(0, m.start() - 40):m.end() + 40].strip()} » [{ts(i)}]")
    return "\n".join([
        "> Résumé extractif de secours : aucun modèle de langage n'a été utilisé. Les phrases sont reprises du transcript.",
        "", "## Sujet", sujet[:200], "", "## Points clés",
        *([f"- {t} [{ts(i)}]" for _, i, t in choisis] or ["- non précisé"]), "",
        "## Personnes, organisations, lieux cités",
        *([f"- {n} [{ts(i)}]" for n, i in list(noms.items())[:20]] or ["- non précisé"]), "",
        "## Chiffres", *(chiffres[:15] or ["- aucun chiffre cité"])])


# ---------------------------------------------------------------- d. sorties
def construire_transcript(meta: dict, lignes: list[dict], inclure_arabe: bool, methode_diar: str) -> str:
    titre = meta.get("titre") or meta.get("sid") or "Session"
    out = [f"# Transcript : {titre}", "",
           f"*Source : {meta.get('source', '?')} · moteur : {meta.get('moteur', '?')} · durée : {horodatage(meta.get('duree_s') or 0)} · "
           f"locuteurs : {methode_diar}*", ""]
    for l in lignes:
        out.append(f"**[{horodatage(l['debut_s'])}] Locuteur {l['locuteur'] or '?'}** : {l['fr_revise'] or '(vide)'}")
        if inclure_arabe and l.get("ar"):
            out.append(f"> {l['ar']}")
        out.append("")
    return "\n".join(out)


def construire_srt(lignes: list[dict], largeur: int = 84) -> str:
    cues, n, prec = [], 1, None
    for l in lignes:
        texte = l["fr_revise"]
        if not texte:
            continue
        if l["locuteur"] != prec and l["locuteur"] is not None:
            texte = f"[Locuteur {l['locuteur']}] {texte}"
        prec = l["locuteur"]
        mots, morceaux, cur = texte.split(), [], ""
        for m in mots:
            if cur and len(cur) + 1 + len(m) > largeur:
                morceaux.append(cur)
                cur = m
            else:
                cur = f"{cur} {m}".strip()
        morceaux.append(cur)
        d, f = float(l["debut_s"]), max(float(l["fin_s"]), float(l["debut_s"]) + 0.5)
        poids = [len(x) for x in morceaux]
        t = d
        for x, p in zip(morceaux, poids):
            fin = t + (f - d) * p / sum(poids)
            cues.append(f"{n}\n{_ts_srt(t)} --> {_ts_srt(fin)}\n{x}\n")
            n += 1
            t = fin
    return "\n".join(cues)


def ecrire_docx(chemin: Path, titre: str, lignes: list[dict], resume_md: str) -> bool:
    try:
        import docx
    except ImportError:
        return False
    d = docx.Document()
    d.add_heading(titre, 0)
    d.add_heading("Résumé", 1)
    for l in resume_md.splitlines():
        if l.startswith("## "):
            d.add_heading(l[3:], 2)
        elif l.startswith("- "):
            d.add_paragraph(l[2:], style="List Bullet")
        elif l.strip() and not l.startswith(">"):
            d.add_paragraph(l)
    d.add_heading("Transcript", 1)
    for l in lignes:
        p = d.add_paragraph()
        p.add_run(f"[{horodatage(l['debut_s'])}] Locuteur {l['locuteur'] or '?'} : ").bold = True
        p.add_run(l["fr_revise"])
    d.save(str(chemin))
    return True


# ---------------------------------------------------------------- orchestration
def generer_rapport(session: Session, options: dict | None = None) -> dict:
    """Produit transcript, résumé, sous-titres et rapport.json dans session.dossier_rapport ; renvoie le rapport.

    options : voir OPTIONS_DEFAUT ; en plus « llm » (callable(messages, max_new_tokens) -> str), « diariseur »,
    « lexique » (liste déjà chargée), « pool » (Pool de lazy_engines à décharger), « progress » (callable(fraction, texte)).
    """
    opt = {**OPTIONS_DEFAUT, **(options or {})}
    prog = opt.get("progress") or (lambda f, d="": None)
    t0 = time.time()
    meta = session.meta()
    segments = session.segments()
    if not segments:
        raise ValueError("la session ne contient aucun segment : rien à rapporter")
    segments.sort(key=lambda s: (s.get("debut_s") or 0))
    avert: list[str] = []
    lexique = opt["lexique"] if opt.get("lexique") is not None else charger_lexique()
    if not lexique:
        avert.append("lexique de noms propres introuvable ou vide : noms non uniformisés")
    sans_fr = not any(texte_de(s) for s in segments)
    if sans_fr:                                          # session en transcription arabe seule
        avert.append("aucune traduction française dans la session : le texte arabe sert de base")
        for s in segments:
            s["fr"] = s.get("ar")

    with verrou_gpu():
        liberer_gpu(opt.get("pool"))
        # a. diarisation
        prog(0.05, "Identification des locuteurs")
        locuteurs, methode_diar = etape_diarisation(session, segments, opt, avert)
        # b. révision
        llm, llm_proprietaire, methode_rev = opt.get("llm"), False, "aucune (désactivée)"
        veut_llm = (opt["revision"] or opt["resume"]) and os.environ.get("PORTEPAROLE_LLM_RAPPORT", "").lower() != "off"
        if llm is None and veut_llm:
            try:
                llm = LLMLocal().charger()
                llm_proprietaire = True
            except Exception as exc:                     # noqa: BLE001 : pas de GPU, mémoire, accès au modèle…
                avert.append(f"LLM local indisponible ({type(exc).__name__}: {str(exc)[:150]}) : pas de relecture, résumé extractif")
                llm = None
        try:
            if opt["revision"] and llm is not None and not sans_fr:
                prog(0.25, "Relecture de la traduction")
                revision, bilan_lex = etape_revision(segments, llm, lexique, opt, lambda f, d: prog(0.25 + 0.5 * f, d))
                methode_rev = f"LLM {getattr(llm, 'nom', 'fourni')}"
            else:
                revision = [{"fr_original": texte_de(s), "fr_revise": texte_de(s), "modifie": False, "raison": "non relu",
                             "entites": [e["fr"] for e in entites_du_texte(s.get("ar") or "", lexique)],
                             "entites_absentes_du_texte": []} for s in segments]
                bilan_lex = {"entites_detectees": sorted({e for r in revision for e in r["entites"]}), "a_verifier": []}
                bilan_lex["n_entites"] = len(bilan_lex["entites_detectees"])
            lignes = [{"id": s.get("id"), "debut_s": s.get("debut_s") or 0, "fin_s": s.get("fin_s") or 0, "locuteur": l,
                       "ar": s.get("ar"), **r} for s, l, r in zip(segments, locuteurs, revision)]
            # c. résumé
            textes = [l["fr_revise"] for l in lignes]
            resume_md, methode_res, controle = "", "aucun", {"cites": 0, "invalides": []}
            if opt["resume"]:
                prog(0.8, "Rédaction du résumé")
                if llm is not None:
                    try:
                        resume_md = resume_llm(lignes_transcript(segments, textes, locuteurs), llm)
                        controle = verifier_horodatages(resume_md, segments)
                        methode_res = "LLM"
                        if controle["invalides"]:
                            avert.append(f"horodatages du résumé absents du transcript : {', '.join(controle['invalides'])}")
                        if not controle["cites"]:
                            avert.append("le résumé ne cite aucun horodatage")
                    except Exception as exc:             # noqa: BLE001
                        avert.append(f"résumé LLM échoué ({type(exc).__name__}) : résumé extractif")
                        resume_md = ""
                if not resume_md:
                    resume_md = resume_extractif(segments, textes, bilan_lex["entites_detectees"])
                    methode_res = "extractif (sans modèle de langage)"
        finally:
            if llm_proprietaire:
                llm.decharger()

    # d. écriture
    prog(0.95, "Écriture des fichiers")
    dossier = session.dossier_rapport
    (dossier / "transcript.md").write_text(construire_transcript(meta, lignes, opt["inclure_arabe"], methode_diar), encoding="utf-8")
    (dossier / "resume.md").write_text(
        f"# Résumé : {meta.get('titre') or meta.get('sid')}\n\n*Méthode : {methode_res}*\n\n{resume_md}\n", encoding="utf-8")
    (dossier / "sous_titres.srt").write_text(construire_srt(lignes), encoding="utf-8")
    fichiers = ["transcript.md", "resume.md", "sous_titres.srt", "rapport.json"]
    if opt["docx"] and ecrire_docx(dossier / "transcript.docx", str(meta.get("titre") or meta.get("sid")), lignes, resume_md):
        fichiers.append("transcript.docx")
    rapport = {
        "sid": session.sid, "genere_le": time.strftime("%Y-%m-%dT%H:%M:%S"), "duree_calcul_s": round(time.time() - t0, 1),
        "methode_diarisation": methode_diar, "methode_revision": methode_rev, "methode_resume": methode_res,
        "n_locuteurs": len({l["locuteur"] for l in lignes if l["locuteur"]}),
        "lexique": bilan_lex, "controle_horodatages": controle, "avertissements": avert,
        "segments": lignes, "resume_markdown": resume_md, "fichiers": fichiers,
    }
    (dossier / "rapport.json").write_text(json.dumps(rapport, ensure_ascii=False, indent=2), encoding="utf-8")
    prog(1.0, "Rapport prêt")
    return rapport
